#!/usr/bin/env python3
"""Build PLOS/Figshare Farrow 30-case veterinarian-consensus urgency task.
Original 30 clinical vignettes + owner scoring are CC BY (cite Farrow et al. 2026).
The author's 2.SCORE compares owner response to a vet-panel consensus; the
unique 'Same urgency assessment' response per vignette recovers the panel's
published 6-level urgency reference. Does NOT recover the 14 individual vet votes.
No respondent-level data or identifiers are exported.
"""
import collections,hashlib,io,json,re,requests,zipfile
from openpyxl import load_workbook
from xml.etree import ElementTree as ET
from pathlib import Path

DOC_URL="https://ndownloader.figshare.com/files/61100344"
XLS_URL="https://ndownloader.figshare.com/files/60375362"
PINNED_DOC_MD5="677ba68a3926088d6ac4da52f3010994"
PINNED_XLS_MD5="f7816195a5e6c5a379862988cf0d1522"
URGENCY=[
 "seek veterinary advice and care immediately (within an hour)",
 "seek veterinary advice and care within 12 hours",
 "seek veterinary advice and care within 24 hours",
 "seek veterinary advice and care within a week (7 days)",
 "seek veterinary advice and care within a fortnight (two weeks - 14 days)",
 "self-care (don't seek veterinary advice and care immediately, try to manage the condition at home)"
]
SHORT=["within_1h","within_12h","within_24h","within_7d","within_14d","self_care"]

def download(url,checksum):
 r=requests.get(url,timeout=45,headers={"User-Agent":"OpenFARM-Research/1.0"});r.raise_for_status()
 if hashlib.md5(r.content).hexdigest()!=checksum:raise RuntimeError("source md5 mismatch")
 return r.content

def main():
 doc=download(DOC_URL,PINNED_DOC_MD5);xlsx=download(XLS_URL,PINNED_XLS_MD5)
 root=ET.fromstring(zipfile.ZipFile(io.BytesIO(doc)).read("word/document.xml"))
 q="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
 tables=root.findall(".//"+q+"tbl")
 if len(tables)<2:raise RuntimeError("no original clinical vignettes table")
 def txt(cell):return "".join((n.text or "") for n in cell.iter(q+"t")).strip()
 source={}
 for row in tables[1].findall("./"+q+"tr")[1:]:
  cells=row.findall("./"+q+"tc")
  if len(cells)<6:raise RuntimeError("malformed S3 table")
  id=txt(cells[0]);signalment=txt(cells[4]);story=txt(cells[5])
  if not id or not signalment or not story or id in source:raise RuntimeError("incomplete or duplicate original vignette")
  source[id]={"id":id,"state":"Dog: "+signalment+"\nPresenting concerns: "+story}
 if len(source)!=30:raise RuntimeError("wrong number of independent scenarios")
 book=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
 ws=book.worksheets[0];it=ws.iter_rows(values_only=True);heads=next(it)
 headers={str(x).strip():j for j,x in enumerate(heads)}
 for k in ("VIGNETTE","2.URGENCY","2.SCORE","2.BinaryScore"):
  if k not in headers:raise RuntimeError("owner survey schema missing "+k)
 scores=collections.defaultdict(lambda:collections.Counter())
 valid=collections.Counter();missing=collections.Counter()
 for row in it:
  id=str(row[headers["VIGNETTE"]]).strip()
  if id not in source:raise RuntimeError("orphan owner vignette")
  urgency=str(row[headers["2.URGENCY"]]).strip()
  score=str(row[headers["2.SCORE"]]).strip()
  binary=str(row[headers["2.BinaryScore"]]).strip()
  if score=="999" or urgency=="999" or binary=="999":
   missing[id]+=1;continue
  if urgency not in URGENCY or score not in {"More urgent","Same urgency assessment","Less urgent"}:
   raise RuntimeError("unknown original response/score")
  if binary!=("Less urgent" if score=="Less urgent" else "Equal to or more urgent"):
   raise RuntimeError("binary comparison mismatch")
  scores[id][(urgency,score)]+=1;valid[id]+=1
 out=[];score_ct=collections.Counter();cats=collections.Counter()
 for id in sorted(source):
  matches=[u for (u,s),n in scores[id].items() if s=="Same urgency assessment"]
  if len(matches)!=1:raise RuntimeError("vet consensus cannot be uniquely recovered")
  consensus=matches[0];index=URGENCY.index(consensus)
  for (urgency,score),n in scores[id].items():
   j=URGENCY.index(urgency)
   expected="More urgent" if j<index else "Less urgent" if j>index else "Same urgency assessment"
   if score!=expected:raise RuntimeError("panel-consensus ordinality not consistent")
   score_ct[score]+=n
  cats[SHORT[index]]+=1
  out.append(dict(**source[id],reference=SHORT[index],
          reference_source="derived_veterinarian_consensus_from_authors_original_2.SCORE",
          n_valid_owner_responses=valid[id],n_missing_owner_responses=missing[id]))
 if sum(valid.values())+sum(missing.values())!=5316:raise RuntimeError("owner response count wrong")
 if score_ct["Less urgent"]!=1497:raise RuntimeError("published 1497 underestimates not replicated")
 if sum(missing.values())!=48:raise RuntimeError("unexpected missing data count")
 folder=Path("reports/index_l3_farrow");folder.mkdir(parents=True,exist_ok=True)
 (folder/"cases_local_reference.json").write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n")
 audit={"vignettes":30,"owner_responses":5316,"valid_owner_responses":sum(valid.values()),
   "missing_owner_responses":sum(missing.values()),"less_urgent":1497,
   "consensus_classes":dict(cats),"distinct_reference_values":len(cats),
   "n_observed_independent_vet_votes":0,
   "source_doc_sha256":hashlib.sha256(doc).hexdigest(),
   "source_owner_sha256":hashlib.sha256(xlsx).hexdigest(),
   "schema":"owner first-contact care-seeking, six ordinal choices; 30 independent expert-consensus cases",
   "rights":"CC BY 4.0; Farrow et al. (2026) PLOS ONE 10.1371/journal.pone.0339723",
   "limitations":"Only consensus recoverable; 14 individual vet rating vectors not provided. Not live triage or patient outcomes."}
 (folder/"source_audit.json").write_text(json.dumps(audit,indent=2)+"\n")
 print("FARROW_BUILDER_RESULT",json.dumps(audit))
if __name__=="__main__":main()
