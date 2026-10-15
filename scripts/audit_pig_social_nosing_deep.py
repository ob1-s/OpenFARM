"""Inspect pig-social-nosing expert ethogram source identities, codes and duplicates.

CC0 source Figshare 27960810 version 2. Reports aggregates only. Not a
video-perception task: source provides manually recorded events but no video.
"""
from __future__ import annotations
import argparse,collections,hashlib,io,json
from pathlib import Path
import pandas as pd
import requests
from openpyxl import load_workbook

ARTICLE=27960810
S=requests.Session()
S.headers.update({"User-Agent":"OpenFARM-CC0-source-ethogram-audit/0.1"})

def get(url):
    r=S.get(url,timeout=(20,45))
    r.raise_for_status()
    if len(r.content)>2_000_000:raise RuntimeError("source cap exceeded")
    return r.content

def file_bytes(name):
    meta=json.loads(get(f"https://api.figshare.com/v2/articles/{ARTICLE}"))
    found=[f for f in meta["files"] if f["name"]==name]
    if len(found)!=1:raise RuntimeError("Expected exactly one named file")
    return get(found[0]["download_url"]),meta

def diagnostics():
    b,meta=file_bytes("All Pen Behaviours.xlsx")
    a,_=file_bytes("All Pen Attributes.xlsx")
    ev=pd.read_excel(io.BytesIO(b))
    animals=pd.read_excel(io.BytesIO(a))
    required=["Batch","Pen","Day","Time_interval","Initiator","Recipient","Behaviour"]
    if list(ev.columns)!=required:raise RuntimeError("Unexpected event table schema")
    if not {"Batch","Pig_id","Finisher_Pen after regrouping"}.issubset(animals.columns):
        raise RuntimeError("Missing independent animal group attributes")
    # Inspect workbook header comments to recover authoritative code semantics.
    comments={}
    book=load_workbook(io.BytesIO(b),read_only=False,data_only=False)
    for sheet in book:
        for row in sheet.iter_rows(min_row=1,max_row=min(6,sheet.max_row)):
            for cell in row:
                if cell.comment:
                    comments[f"{sheet.title}/{cell.coordinate}"]=cell.comment.text[:900]
        if len(comments)>40:break
    codes=ev["Behaviour"].astype(str)
    norm=codes.str.strip()
    unique_animals=animals[["Batch","Finisher_Pen after regrouping","Pig_id"]].drop_duplicates()
    unique_interaction_ids=pd.concat([
        ev[["Batch","Pen","Initiator"]].rename(columns={"Initiator":"Pig"}),
        ev[["Batch","Pen","Recipient"]].rename(columns={"Recipient":"Pig"})
    ]).drop_duplicates()
    attrkeys=set(tuple(x) for x in unique_animals.itertuples(index=False,name=None))
    # Pen name mapping in two source sheets appears different ("FP1" vs "P1"):
    # audit both literal and canonical suffix, never silently join with guesses.
    eventkeys=set((str(r.Batch),str(r.Pen),str(r.Pig)) for r in unique_interaction_ids.itertuples(index=False))
    orig_keys=set((str(r.Batch),str(r[1]),str(r[2])) for r in unique_animals.itertuples(index=False,name=None))
    clean_keys=set((batch,pen.replace("FP","P"),pig) for batch,pen,pig in orig_keys)
    answer={
        "source":"https://figshare.com/articles/dataset/27960810",
        "doi":meta.get("doi"),"source_license":meta.get("license",{}),
        "revision":meta.get("version"),
        "original_file_sha256":{"behaviour":hashlib.sha256(b).hexdigest(),"attributes":hashlib.sha256(a).hexdigest()},
        "behavior_rows":len(ev),"attribute_rows":len(animals),
        "events_exact_duplicate_rows":int(ev.duplicated().sum()),
        "duplicate_warning":"Same behavior within same coarse interval can happen repeatedly; don't drop duplicates without original observer protocol.",
        "batch_distribution":{str(k):int(v) for k,v in ev.Batch.value_counts().items()},
        "day_distribution":{str(k):int(v) for k,v in ev.Day.value_counts().items()},
        "interval_distribution":{str(k):int(v) for k,v in ev.Time_interval.value_counts().head(24).items()},
        "observed_codes":{str(k):int(v) for k,v in codes.value_counts().items()},
        "normalized_codes":{str(k):int(v) for k,v in norm.value_counts().items()},
        "attribute_unique_batch_pen_pig":len(attrkeys),
        "event_unique_batch_pen_actor":len(eventkeys),
        "literal_join_coverage":len(eventkeys&orig_keys),
        "candidate_pen_renaming_join_coverage":len(eventkeys&clean_keys),
        "workbook_comment_definitions":comments,
        "source_media":"No synchronized event video in Figshare record: structured human ethogram only.",
        "task_status":"Source metadata verified; no claim of supervised video perception or validated positive emotion.",
    }
    return answer

if __name__=="__main__":
    p=argparse.ArgumentParser()
    p.add_argument("--output",type=Path,required=True)
    a=p.parse_args()
    result=diagnostics()
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(result,indent=2,default=str)+"\n")
    print("PIG_AFFILIATIVE_DEEP_AUDIT:"+json.dumps(result,default=str))
