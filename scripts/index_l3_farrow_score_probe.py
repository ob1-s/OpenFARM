#!/usr/bin/env python3
"""Aggregate-only audit of Farrow published owner responses and S1 vignettes.
No respondent identifiers, survey rows or individual responses are logged/published.
"""
import collections, io, json, re, requests, zipfile
from openpyxl import load_workbook
from xml.etree import ElementTree as ET
DOC_URL="https://ndownloader.figshare.com/files/61100344"
XLS_URL="https://ndownloader.figshare.com/files/60375362"
S=requests.Session();S.headers.update({"User-Agent":"OpenFARM-Research/1.0"})
doc=S.get(DOC_URL,timeout=45);doc.raise_for_status()
data=S.get(XLS_URL,timeout=45);data.raise_for_status()
w=load_workbook(io.BytesIO(data.content),read_only=True,data_only=True)
ws=w.worksheets[0];rows=ws.iter_rows(values_only=True);header=next(rows);idx={str(k).strip():i for i,k in enumerate(header)}
KEYS=["VIGNETTE","2.URGENCY","2.SCORE","2.BinaryScore"]
for k in KEYS:
    if k not in idx:raise RuntimeError(f"missing official column {k}")
by=collections.defaultdict(lambda:collections.defaultdict(collections.Counter))
n=0
for row in rows:
    n+=1
    id=str(row[idx["VIGNETTE"]]).strip()
    raw=row[idx["2.URGENCY"]];sc=row[idx["2.SCORE"]];bin_=row[idx["2.BinaryScore"]]
    by[id]["urgency_to_score"][(str(raw),str(sc))]+=1
    by[id]["binary_to_score"][(str(bin_),str(sc))]+=1
    by[id]["triple"][(str(raw),str(sc),str(bin_))]+=1
x=collections.Counter()
summary=[]
for id in sorted(by):
    rec=by[id]
    x.update({k:len(rec[k]) for k in rec})
    summary.append({"vignette":id,"n":sum(rec["triple"].values()),
                    "urgency_by_score":[{"urgency":a,"score":b,"n":v} for (a,b),v in sorted(rec["urgency_to_score"].items())],
                    "binary_by_score":[{"binary":a,"score":b,"n":v} for (a,b),v in sorted(rec["binary_to_score"].items())]})
# Source DOCX: safe schema only, no individual clinical vignettes logged.
with zipfile.ZipFile(io.BytesIO(doc.content)) as z:
    root=ET.fromstring(z.read("word/document.xml"))
    q="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    tables=root.findall(".//"+q+"tbl")
    names=[]
    table=tables[1]
    for row in table.findall("./"+q+"tr")[1:]:
        cells=row.findall("./"+q+"tc")
        def txt(cell):return "".join((n.text or "") for n in cell.iter(q+"t")).strip()
        if len(cells)>5:
            names.append(txt(cells[0]))
result={"n_owner_responses":n,"n_distinct_vignettes":len(by),
        "published_doc_vignette_ids":names,"per_vignette":summary,
        "interpretation":"Only original owners' urgency answers and author-coded comparison variables; reference category must be uniquely inferred from equal-to-vet scoring or original vet data."}
import os
os.makedirs("reports",exist_ok=True)
with open("reports/index_l3_farrow_gold_score_audit.json","w") as f:json.dump(result,f,indent=2)
print("farrow n_owner_responses",n,"n_distinct_vignettes",len(by),"doc",len(names))
for r in summary:
    print("FARROW_AGG",json.dumps(r,ensure_ascii=False)[:3200])
